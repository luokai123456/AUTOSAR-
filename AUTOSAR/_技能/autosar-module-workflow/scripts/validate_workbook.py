# -*- coding: utf-8 -*-
"""通用校验：<模块名>标准化流程.xlsx 的结构、日期、ID、枚举、隐私、数据验证、甘特视图。"""
import argparse
import datetime as dt
import re
import sys
from openpyxl import load_workbook

EXPECTED_SHEETS = [
    "0-使用说明", "0a-目录说明", "1-文件变更履历", "2-流程总览", "3-开发计划",
    "4-需求追溯矩阵", "5-差异台账", "6-测试用例与CAPL", "7-问题台账",
    "8-评审记录", "9-变更记录", "10-基线清单", "11-能力分析", "12-术语与参考", "13-甘特视图",
]
STATUS = {"计划", "进行中", "待评审", "完成", "阻塞", "挂起"}
APPLY = {"适用", "裁剪", "差异", "待确认"}
DIFF_TYPES = {"可选特性未实现", "实现偏差", "OEM特殊需求", "工具链限制"}
GATE = {"通过", "有条件通过", "不通过", "待评审"}
CASE_STATUS = {"计划", "执行中", "通过", "失败", "阻塞"}
ISSUE_STATUS = {"已解决", "规避中", "待确认", "无法复现"}
PHONE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")

errors, warnings = [], []


def check(cond, msg):
    if not cond:
        errors.append(msg)


def is_date(v):
    return isinstance(v, (dt.date, dt.datetime))


def check_ids(ws, col, prefix, label):
    seen = set()
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, col).value
        if v is None:
            continue
        s = str(v).strip()
        check(s.startswith(prefix), "%s 第%d行 ID 前缀错误: %s" % (label, r, s))
        check(s not in seen, "%s 第%d行 ID 重复: %s" % (label, r, s))
        seen.add(s)
    return len(seen)


def main(workbook, module):
    wb = load_workbook(workbook)
    check(wb.sheetnames == EXPECTED_SHEETS, "Sheet 列表不符: %r" % (wb.sheetnames,))

    # 3-开发计划
    ws = wb["3-开发计划"]
    headers = [ws.cell(2, c).value for c in range(1, 18)]
    need = ["阶段", "WBS", "任务项", "交付物ID", "交付物", "验收标准", "前置依赖", "工期(天)",
            "计划开始", "计划完成", "实际开始", "实际完成", "进度", "状态", "责任人",
            "关联需求/差异", "风险与备注"]
    check(headers == need, "开发计划表头不符: %r" % (headers,))
    plan_rows = 0
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 2).value is None:
            continue
        plan_rows += 1
        check(ws.cell(r, 1).value in {"P0", "P1", "P2", "P3", "P4", "P5", "P6"},
              "计划表第%d行阶段非法: %r" % (r, ws.cell(r, 1).value))
        for c in (9, 10, 11, 12):
            v = ws.cell(r, c).value
            check(v is None or is_date(v), "计划表第%d行日期非日期类型: %r" % (r, v))
        p = ws.cell(r, 13).value
        check(p is None or (isinstance(p, (int, float)) and 0 <= p <= 1), "计划表第%d行进度越界: %r" % (r, p))
        check(ws.cell(r, 14).value in STATUS, "计划表第%d行状态非法: %r" % (r, ws.cell(r, 14).value))
        check(ws.cell(r, 5).value and ws.cell(r, 6).value, "计划表第%d行缺交付物/验收标准" % r)
    check(plan_rows >= 20, "开发计划任务数偏少: %d" % plan_rows)

    # 4-需求追溯
    ws = wb["4-需求追溯矩阵"]
    trace_rows = 0
    seen_req = set()
    prefix = "SWS_%s_" % module
    for r in range(3, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if rid is None:
            continue
        trace_rows += 1
        check(str(rid).startswith(prefix), "需求追溯第%d行需求号前缀应为 %s: %r" % (r, prefix, rid))
        check(str(rid) not in seen_req, "需求追溯需求号重复: %s" % rid)
        seen_req.add(str(rid))
        check(ws.cell(r, 4).value in APPLY, "需求追溯第%d行适用性非法: %r" % (r, ws.cell(r, 4).value))
        if ws.cell(r, 4).value in {"裁剪", "差异"}:
            check(ws.cell(r, 9).value, "需求追溯第%d行裁剪/差异缺备注依据" % r)

    # 5-差异
    ws = wb["5-差异台账"]
    n_diff = check_ids(ws, 1, "DIF-", "差异台账")
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value:
            check(ws.cell(r, 2).value in DIFF_TYPES, "差异台账第%d行类型非法: %r" % (r, ws.cell(r, 2).value))
            check(ws.cell(r, 7).value, "差异台账第%d行缺处置结论" % r)

    # 6-用例
    ws = wb["6-测试用例与CAPL"]
    n_case = check_ids(ws, 1, "TC-", "测试用例")
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value:
            check(ws.cell(r, 9).value in CASE_STATUS, "测试用例第%d行状态非法: %r" % (r, ws.cell(r, 9).value))
            check(ws.cell(r, 6).value, "测试用例第%d行缺关联需求" % r)

    # 7-问题
    ws = wb["7-问题台账"]
    n_issue = check_ids(ws, 1, "Q-", "问题台账")
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value:
            check(ws.cell(r, 7).value in ISSUE_STATUS, "问题台账第%d行状态非法: %r" % (r, ws.cell(r, 7).value))

    # 8-评审
    ws = wb["8-评审记录"]
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value:
            check(re.fullmatch(r"G[0-6]", str(ws.cell(r, 1).value)), "评审记录第%d行评审门非法" % r)
            check(ws.cell(r, 5).value in GATE, "评审记录第%d行结论非法: %r" % (r, ws.cell(r, 5).value))

    # 10-基线
    ws = wb["10-基线清单"]
    n_base = check_ids(ws, 1, "BL-", "基线清单")

    # 日期类型
    for name, col in [("7-问题台账", 2), ("8-评审记录", 2), ("9-变更记录", 2),
                      ("10-基线清单", 2), ("1-文件变更履历", 1)]:
        ws = wb[name]
        for r in range(3, ws.max_row + 1):
            v = ws.cell(r, col).value
            check(v is None or is_date(v), "%s 第%d行日期非日期类型: %r" % (name, r, v))

    # 数据验证
    for name in ["3-开发计划", "4-需求追溯矩阵", "5-差异台账", "6-测试用例与CAPL",
                 "7-问题台账", "8-评审记录"]:
        check(len(wb[name].data_validations.dataValidation) >= 1, "%s 缺少数据验证" % name)

    # 甘特视图
    check("13-甘特视图" in wb.sheetnames, "缺少 13-甘特视图")
    if "13-甘特视图" in wb.sheetnames:
        g = wb["13-甘特视图"]
        check(g.max_column >= 15, "甘特视图列数不足: %d" % g.max_column)
        check(g.max_row >= 10, "甘特视图行数不足: %d" % g.max_row)
        title = str(g.cell(1, 1).value or "")
        check(module in title, "甘特标题未包含模块名: %r" % title)
        fills = 0
        for row in g.iter_rows(min_row=3, max_row=g.max_row):
            for c in row:
                try:
                    if c.fill and c.fill.fgColor and c.fill.fgColor.rgb not in (None, "00000000"):
                        fills += 1
                except Exception:
                    pass
        check(fills > 20, "甘特视图着色单元格过少: %d" % fills)

    # 公式错误 / 隐私
    for s in wb.worksheets:
        for row in s.iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                if "#REF!" in c.value:
                    errors.append("%s!%s 含 #REF!" % (s.title, c.coordinate))
                if PHONE.search(c.value):
                    errors.append("%s!%s 含手机号" % (s.title, c.coordinate))
                if EMAIL.search(c.value):
                    errors.append("%s!%s 含邮箱" % (s.title, c.coordinate))

    print("=" * 64)
    print("校验文件:", workbook, "| 模块:", module)
    print("Sheet 数: %d | 计划任务: %d | 需求追溯: %d | 差异: %d | 用例: %d | 问题: %d | 基线: %d"
          % (len(wb.sheetnames), plan_rows, trace_rows, n_diff, n_case, n_issue, n_base))
    print("错误: %d | 警告: %d" % (len(errors), len(warnings)))
    for e in errors[:50]:
        print("  [FAIL]", e)
    print("结果:", "PASS" if not errors else "FAIL")
    return 0 if not errors else 1


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="校验模块标准化流程工作簿")
    ap.add_argument("workbook")
    ap.add_argument("--module", default="CanNm")
    a = ap.parse_args()
    sys.exit(main(a.workbook, a.module))
