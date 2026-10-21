# -*- coding: utf-8 -*-
"""由现有模块工作簿实例化新模块工作簿（EcuM/Com 等）。"""
import argparse
import datetime as dt
import os
import sys

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from refresh_gantt import build_gantt  # noqa: E402

THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def replace_tokens(wb, module):
    repl = [("CanNm", module), ("CANNM", module)]
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    v = c.value
                    for a, b in repl:
                        v = v.replace(a, b)
                    if v != c.value:
                        c.value = v
                        n += 1
    return n


def clear_rows(ws, start=3):
    if ws.max_row >= start:
        ws.delete_rows(start, ws.max_row - start + 1)


def fill_rows(ws, rows):
    r = 3
    for row in rows:
        for j, v in enumerate(row, 1):
            c = ws.cell(r, j, v)
            c.border = BORDER
            c.font = Font(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1


def instantiate(source, target, module):
    wb = load_workbook(source)

    # 1) 模块名替换
    replace_tokens(wb, module)

    # 2) 清空模块特定内容，写入模板行
    ws = wb["4-需求追溯矩阵"]
    clear_rows(ws)
    fill_rows(ws, [
        ("SWS_%s_00001" % module, "7.x", "（待整理：按 SWS 需求号填写中文摘要）", "适用",
         "（待填写：调用逻辑/函数）", "（待填写：ECUC 参数）", "TC-001", "计划", "示例行，替换为真实需求"),
    ])

    ws = wb["5-差异台账"]
    clear_rows(ws)
    fill_rows(ws, [
        ("DIF-001", "可选特性未实现", "（SWS 章节/可选特性）", "（差异描述）", "（来源）",
         "（影响分析）", "（处置：实现/规避/接受/CR）", "（责任人）", "待确认", "-", "（验证方式）"),
    ])

    ws = wb["6-测试用例与CAPL"]
    clear_rows(ws)
    fill_rows(ws, [
        ("TC-001", "（场景）", "（前置条件）", "（步骤）", "（预期结果）",
         "（关联需求）", "CAPL_001.can", "", "计划"),
    ])

    ws = wb["7-问题台账"]
    clear_rows(ws)
    fill_rows(ws, [
        ("Q-001", dt.date.today(), "（现象）", "（环境/版本）", "（根因）", "（解决/规避）", "待确认", "（关联文档）"),
    ])

    ws = wb["8-评审记录"]
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value:
            ws.cell(r, 2).value = None
            ws.cell(r, 5).value = "待评审"
            ws.cell(r, 6).value = ""

    ws = wb["10-基线清单"]
    clear_rows(ws)
    fill_rows(ws, [
        ("BL-001", dt.date.today(), "R21-11 CP", "（工具版本）", "（BSWMD/ECUC 版本）", "-", "V0.01", "-", "实例化基线"),
        ("BL-002", None, "R21-11 CP", "", "", "", "", "", "P6.2 归档"),
    ])

    # 3) 变更履历：改为实例化记录
    ws = wb["1-文件变更履历"]
    clear_rows(ws)
    fill_rows(ws, [
        (dt.date.today(), "V0.01", "由《CANNM标准化流程.xlsx》实例化为 %s 模块流程（结构继承，内容待填）" % module, "（责任人）"),
    ])

    # 4) 能力分析/术语中的模块名已替换；刷新甘特
    build_gantt(wb, module)
    wb.save(target)
    print("instantiated:", target, "module:", module)
    print("next: 填写 3-开发计划 日期/基线 → 刷新甘特 → 运行 validate_workbook.py")


def main():
    ap = argparse.ArgumentParser(description="实例化模块标准化流程工作簿")
    ap.add_argument("--source", required=True)
    ap.add_argument("--target", required=True)
    ap.add_argument("--module", required=True)
    args = ap.parse_args()
    instantiate(args.source, args.target, args.module)


if __name__ == "__main__":
    main()
