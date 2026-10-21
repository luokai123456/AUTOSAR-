# -*- coding: utf-8 -*-
"""检查 CANNM学习计划.xlsx 的结构与内容"""
import os, sys
from openpyxl import load_workbook

P = r"C:\Users\11401\Desktop\学习\AUTOSAR\00_学习路线\CANNM学习计划.xlsx"
out = []
wb = load_workbook(P, data_only=False)
out.append("sheets: %r" % (wb.sheetnames,))
for ws in wb.worksheets:
    out.append("\n===== SHEET: %s | dims=%s | merged=%d | images=%d | charts=%d | tables=%d" % (
        ws.title, ws.dimensions, len(ws.merged_cells.ranges), len(ws._images), len(ws._charts), len(ws.tables)))
    out.append("merged: %r" % ([str(r) for r in ws.merged_cells.ranges][:20],))
    out.append("freeze: %r" % (ws.freeze_panes,))
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 120), max_col=min(ws.max_column, 20)))
    for r in rows:
        vals = []
        for c in r:
            v = c.value
            if v is None:
                vals.append("")
            else:
                s = str(v).replace("\n", " / ")
                vals.append(s[:120])
        if any(vals):
            out.append("R%02d | " % r[0].row + " | ".join(vals))
open(r"C:\Users\11401\Desktop\学习\.tools\plan_inspect.txt", "w", encoding="utf-8").write("\n".join(out))
print("done, sheets:", wb.sheetnames)
