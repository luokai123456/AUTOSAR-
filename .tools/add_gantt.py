# -*- coding: utf-8 -*-
"""为 CANNM标准化流程.xlsx 增加「13-甘特视图」（按周粒度，从 3-开发计划 生成）。"""
import datetime as dt
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

P = r"C:\Users\11401\Desktop\学习\AUTOSAR\00_学习路线\CANNM标准化流程.xlsx"
wb = load_workbook(P)

if "13-甘特视图" in wb.sheetnames:
    del wb["13-甘特视图"]
ws = wb.create_sheet("13-甘特视图")
ws.sheet_view.showGridLines = False

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PHASE = {
    "P0": "1F4E79", "P1": "2E75B6", "P2": "70AD47", "P3": "FFC000",
    "P4": "ED7D31", "P5": "7030A0", "P6": "808080",
}
GATE = "FFD966"

plan = wb["3-开发计划"]
tasks = []
for r in range(3, plan.max_row + 1):
    wbs = plan.cell(r, 2).value
    if wbs is None:
        continue
    start, end = plan.cell(r, 9).value, plan.cell(r, 10).value
    if not (isinstance(start, (dt.date, dt.datetime)) and isinstance(end, (dt.date, dt.datetime))):
        continue
    tasks.append({
        "stage": plan.cell(r, 1).value, "wbs": str(wbs), "name": plan.cell(r, 3).value,
        "start": start.date() if isinstance(start, dt.datetime) else start,
        "end": end.date() if isinstance(end, dt.datetime) else end,
        "days": plan.cell(r, 8).value, "progress": plan.cell(r, 13).value,
    })

min_d = min(t["start"] for t in tasks) if tasks else dt.date(2026, 10, 21)
max_d = max(t["end"] for t in tasks) if tasks else dt.date(2027, 2, 1)
# 周一对齐
week0 = min_d - dt.timedelta(days=min_d.weekday())
weeks = []
d = week0
while d <= max_d:
    weeks.append(d)
    d += dt.timedelta(days=7)

FIXED = ["WBS", "任务项", "阶段", "计划开始", "计划完成", "工期(天)", "进度"]
n_fixed = len(FIXED)
last_col = n_fixed + len(weeks)

# 标题
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
c = ws.cell(1, 1, "CanNm 标准化流程甘特视图（按周）")
c.font = Font(bold=True, size=14, color="FFFFFF")
c.fill = HEAD_FILL
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 26

# 表头
for j, h in enumerate(FIXED, 1):
    cell = ws.cell(2, j, h)
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
for k, w in enumerate(weeks):
    col = n_fixed + 1 + k
    cell = ws.cell(2, col, "%d/%d" % (w.month, w.day))
    cell.font = Font(bold=True, color="FFFFFF", size=8)
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER

# 任务行
row = 3
for t in tasks:
    ws.cell(row, 1, t["wbs"])
    ws.cell(row, 2, t["name"])
    ws.cell(row, 3, t["stage"])
    ws.cell(row, 4, t["start"]).number_format = "yyyy-mm-dd"
    ws.cell(row, 5, t["end"]).number_format = "yyyy-mm-dd"
    ws.cell(row, 6, t["days"])
    pv = ws.cell(row, 7, t["progress"])
    pv.number_format = "0%"
    for j in range(1, n_fixed + 1):
        ws.cell(row, j).border = BORDER
        ws.cell(row, j).font = Font(size=9)
        if j in (1, 3, 6, 7):
            ws.cell(row, j).alignment = Alignment(horizontal="center", vertical="center")
        else:
            ws.cell(row, j).alignment = Alignment(vertical="center", wrap_text=(j == 2))
    color = PHASE.get(t["stage"], "A6A6A6")
    is_gate = str(t["wbs"]).startswith("G")
    for k, w in enumerate(weeks):
        col = n_fixed + 1 + k
        week_end = w + dt.timedelta(days=6)
        cell = ws.cell(row, col)
        cell.border = BORDER
        if is_gate and w <= t["end"] <= week_end:
            cell.value = "◆"
            cell.fill = PatternFill("solid", fgColor=GATE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(size=9, bold=True)
        elif w <= t["end"] and week_end >= t["start"]:
            cell.fill = PatternFill("solid", fgColor=color)
    ws.row_dimensions[row].height = 20
    row += 1

# 图例
row += 1
ws.cell(row, 1, "图例：").font = Font(bold=True, size=9)
names = {"P0": "P0 计划与基线", "P1": "P1 标准研读", "P2": "P2 代码与配置映射",
         "P3": "P3 差异分析", "P4": "P4 详设与测试设计", "P5": "P5 测试与执行", "P6": "P6 固化"}
col = 2
for k, v in names.items():
    cell = ws.cell(row, col, v)
    cell.fill = PatternFill("solid", fgColor=PHASE[k])
    cell.font = Font(size=9, color="FFFFFF" if k not in ("P3",) else "000000")
    cell.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    col += 2
cell = ws.cell(row, col, "◆ 评审门")
cell.fill = PatternFill("solid", fgColor=GATE)
cell.font = Font(size=9)
cell.alignment = Alignment(horizontal="center")

row += 1
ws.cell(row, 1, "说明：按周粒度；数据来自「3-开发计划」，修改计划后重新运行 refresh_gantt 脚本刷新。").font = Font(size=9, italic=True, color="808080")

# 列宽
widths = [8, 34, 6, 12, 12, 9, 8] + [4] * len(weeks)
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "H3"
ws.sheet_properties.tabColor = "00B050"

# 更新目录说明
toc = wb["0a-目录说明"]
tr = toc.max_row + 1
vals = ["13-甘特视图", "按周展示计划进度的甘特视图（由 3-开发计划 生成）", "模块负责人"]
for j, v in enumerate(vals, 1):
    c = toc.cell(tr, j, v)
    c.border = BORDER
    c.font = Font(size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(P)
print("gantt added: %d tasks x %d weeks, last_col=%d" % (len(tasks), len(weeks), last_col))
